import os
import io
import json
import re
import logging
import tempfile

from dotenv import load_dotenv
load_dotenv()

from langchain_cohere import CohereEmbeddings
from langchain_groq import ChatGroq
from langchain_community.document_loaders import PyPDFLoader
from langchain_text_splitters import RecursiveCharacterTextSplitter
from langchain_qdrant import Qdrant
from langchain_core.documents import Document

from qdrant_client import QdrantClient
from qdrant_client.http import models
from qdrant_client.http.exceptions import ResponseHandlingException, UnexpectedResponse

import pandas as pd
from PyPDF2 import PdfReader

try:
    import docx  # python-docx
except ImportError:
    docx = None


class SmoothException(Exception):
    """Custom exception class"""
    pass


SUPPORTED_EXTENSIONS = {".pdf", ".docx", ".txt"}


class ResumeRanker:
    def __init__(self, similarity_weight: float = 0.5, llm_weight: float = 0.5):
        """
        similarity_weight + llm_weight should sum to 1.0.
        similarity_weight controls how much the raw embedding/cosine
        similarity (NLP similarity) contributes to the final Score,
        versus the LLM's qualitative judgement.
        """
        self.logger = logging.getLogger(__name__)
        total = similarity_weight + llm_weight
        self.similarity_weight = similarity_weight / total
        self.llm_weight = llm_weight / total

    # ------------------------------------------------------------------
    # Model loading
    # ------------------------------------------------------------------
    def load_models(self, cohere_key, groq_key):
        """Load Cohere embeddings and Groq LLM"""
        embeddings = CohereEmbeddings(cohere_api_key=cohere_key, model="embed-english-v3.0")
        llm = ChatGroq(
            groq_api_key=groq_key,
            model_name="llama-3.3-70b-versatile",
            temperature=0.1
        )
        return embeddings, llm

    # ------------------------------------------------------------------
    # File validation
    # ------------------------------------------------------------------
    def validate_file(self, file_content, filename):
        """Validate a resume file (PDF/DOCX/TXT) before processing"""
        ext = os.path.splitext(filename)[1].lower()

        if ext not in SUPPORTED_EXTENSIONS:
            return False, f"Unsupported file type '{ext}'. Supported: PDF, DOCX, TXT"

        if len(file_content) == 0:
            return False, "File is empty"

        try:
            if ext == ".pdf":
                pdf_file = io.BytesIO(file_content)
                reader = PdfReader(pdf_file)

                if reader.is_encrypted:
                    return False, "PDF is encrypted/password protected"
                if len(reader.pages) == 0:
                    return False, "PDF has no pages"

                text = reader.pages[0].extract_text()
                if not text or len(text.strip()) < 10:
                    return False, "PDF appears to be empty or contains no readable text"

            elif ext == ".docx":
                if docx is None:
                    return False, "python-docx is not installed on the server"
                doc_file = io.BytesIO(file_content)
                d = docx.Document(doc_file)
                text = "\n".join(p.text for p in d.paragraphs)
                if not text or len(text.strip()) < 10:
                    return False, "DOCX appears to be empty or contains no readable text"

            elif ext == ".txt":
                text = file_content.decode("utf-8", errors="ignore")
                if not text or len(text.strip()) < 10:
                    return False, "TXT file appears to be empty"

            return True, "File is valid"

        except Exception as e:
            return False, f"File validation failed: {str(e)}"

    # Backwards-compatible alias
    def validate_pdf(self, file_content, filename):
        return self.validate_file(file_content, filename)

    # ------------------------------------------------------------------
    # Parsing (PDF / DOCX / TXT) -> list[Document]
    # ------------------------------------------------------------------
    def parse_pdf(self, file_content, filename):
        with tempfile.NamedTemporaryFile(delete=False, suffix=".pdf") as tmp:
            tmp.write(file_content)
            tmp_path = tmp.name
        try:
            loader = PyPDFLoader(tmp_path)
            docs = loader.load()
            if not docs:
                raise Exception("No content extracted from PDF")
            for d in docs:
                d.metadata["source"] = filename
                d.metadata["filename"] = filename
            return docs
        except Exception as e:
            raise Exception(f"PDF parsing error: {str(e)}")
        finally:
            if os.path.exists(tmp_path):
                os.unlink(tmp_path)

    def parse_docx(self, file_content, filename):
        if docx is None:
            raise Exception("python-docx is not installed. Run: pip install python-docx")
        try:
            doc_file = io.BytesIO(file_content)
            d = docx.Document(doc_file)
            full_text = "\n".join(p.text for p in d.paragraphs if p.text.strip())

            # Also pull text out of tables (skills tables, etc.)
            for table in d.tables:
                for row in table.rows:
                    row_text = " | ".join(cell.text for cell in row.cells if cell.text.strip())
                    if row_text.strip():
                        full_text += "\n" + row_text

            if not full_text.strip():
                raise Exception("No content extracted from DOCX")

            return [Document(
                page_content=full_text,
                metadata={"source": filename, "filename": filename}
            )]
        except Exception as e:
            raise Exception(f"DOCX parsing error: {str(e)}")

    def parse_txt(self, file_content, filename):
        try:
            text = file_content.decode("utf-8", errors="ignore")
            if not text.strip():
                raise Exception("No content extracted from TXT")
            return [Document(
                page_content=text,
                metadata={"source": filename, "filename": filename}
            )]
        except Exception as e:
            raise Exception(f"TXT parsing error: {str(e)}")

    def parse_resume(self, file_content, filename):
        """Dispatch to the correct parser based on file extension."""
        ext = os.path.splitext(filename)[1].lower()
        if ext == ".pdf":
            return self.parse_pdf(file_content, filename)
        elif ext == ".docx":
            return self.parse_docx(file_content, filename)
        elif ext == ".txt":
            return self.parse_txt(file_content, filename)
        else:
            raise Exception(f"Unsupported file type: {ext}")

    # ------------------------------------------------------------------
    # Vector store
    # ------------------------------------------------------------------
    def create_vector_store(self, docs, collection, qdrant_host, qdrant_api_key, dimension, embeddings):
        """Create/populate vector store for resume chunks"""
        try:
            client = QdrantClient(url=qdrant_host, api_key=qdrant_api_key, timeout=60.0)

            if not client.collection_exists(collection_name=collection):
                client.create_collection(
                    collection_name=collection,
                    vectors_config=models.VectorParams(size=dimension, distance=models.Distance.COSINE),
                )
            else:
                # Clear old points so results from previous runs don't leak in
                client.delete_collection(collection_name=collection)
                client.create_collection(
                    collection_name=collection,
                    vectors_config=models.VectorParams(size=dimension, distance=models.Distance.COSINE),
                )

            vector_store = Qdrant.from_documents(
                documents=docs,
                embedding=embeddings,
                url=qdrant_host,
                api_key=qdrant_api_key,
                collection_name=collection,
            )
            return vector_store
        except (ResponseHandlingException, UnexpectedResponse) as e:
            raise SmoothException(e.reason_phrase)

    def query_resumes(self, query, collection, qdrant_host, qdrant_api_key, embeddings, k):
        """Query resumes based on job description. Returns cosine-similarity scored chunks."""
        try:
            client = QdrantClient(url=qdrant_host, api_key=qdrant_api_key)
            embeddings_data = embeddings.embed_query(query)

            result = client.query_points(
                collection_name=collection,
                query=embeddings_data,
                limit=k,
                with_payload=True
            )

            results = []
            for point in result.points:
                payload = point.payload or {}
                # langchain-qdrant nests page_content/metadata depending on version
                content = payload.get("page_content", "")
                metadata = payload.get("metadata", {}) or {}
                source = (
                    payload.get("source")
                    or payload.get("filename")
                    or metadata.get("source")
                    or metadata.get("filename")
                    or "Unknown"
                )
                results.append({
                    "content": content,
                    "score": point.score,  # cosine similarity, 0-1
                    "source": source
                })
            return results
        except Exception as e:
            self.logger.error(f"Qdrant query error: {str(e)}")
            raise SmoothException(f"Qdrant query failed: {str(e)}")

    # ------------------------------------------------------------------
    # LLM extraction + scoring
    # ------------------------------------------------------------------
    def analyze_resume_match(self, llm, job_desc, resume_text, resume_name):
        """
        Extract structured skills/experience/education AND produce a
        qualitative LLM match score + reasoning, in a single call.
        """
        prompt = f"""
        You are a technical recruiter. Analyze this resume against the job description.

        JOB DESCRIPTION:
        {job_desc[:1500]}

        RESUME CONTENT:
        {resume_text[:3000]}

        Respond ONLY with valid JSON in this EXACT format (no markdown fences, no preamble):
        {{
            "skills": ["Python", "Django", "AWS"],
            "experience": "5 years as a backend engineer, 2 years leading a team of 4",
            "education": "B.Sc. Computer Science, XYZ University",
            "llm_score": 85,
            "strengths": "Strong technical skills, relevant project experience",
            "gaps": "Missing cloud certification, limited team leadership",
            "overall_assessment": "Excellent match with strong alignment to required skills"
        }}

        Scoring guidelines for "llm_score" (0-100):
        - 90-100: Excellent match (most requirements met, strong alignment)
        - 80-89: Very good match (most requirements met, minor gaps)
        - 70-79: Good match (core requirements met, some gaps)
        - 60-69: Partial match (some requirements met, significant gaps)
        - Below 60: Poor match (few requirements met)

        Base "skills", "experience", and "education" strictly on what is stated in the
        resume text above. Do not invent information that isn't present.
        """

        default = {
            "skills": [],
            "experience": "Not specified",
            "education": "Not specified",
            "llm_score": 50,
            "strengths": "Analysis unavailable",
            "gaps": "Analysis unavailable",
            "overall_assessment": "Analysis incomplete"
        }

        try:
            response = llm.invoke(prompt)
            json_match = re.search(r"\{.*\}", response.content, re.DOTALL)
            if json_match:
                analysis = json.loads(json_match.group())
                merged = {**default, **analysis}
                return merged
            return default
        except Exception as e:
            print(f"Analysis error for {resume_name}: {str(e)}")
            default["overall_assessment"] = f"Error in processing: {str(e)}"
            return default

    # ------------------------------------------------------------------
    # Main ranking pipeline
    # ------------------------------------------------------------------
    def rank_resumes(self, job_desc, resume_files_data, cohere_key, groq_key, k=10):
        """
        Main function to rank resumes against a job description.

        resume_files_data: list of {"name": filename, "content": bytes}
        Returns: pandas.DataFrame sorted by final blended Score, descending.
        """
        if not cohere_key or not groq_key:
            raise Exception("Cohere & Groq API Keys are required!")
        if not resume_files_data:
            raise Exception("No resume files provided!")

        QDRANT_HOST = os.getenv(
            "QDRANT_URL",
            "https://923af7ee-921b-41eb-b815-95b55e8ccd55.us-west-1-0.aws.cloud.qdrant.io"
        )
        QDRANT_API_KEY = os.getenv("QDRANT_API_KEY")
        COLLECTION_NAME = "resume_collection"
        DIMENSION = 1024

        embeddings, llm = self.load_models(cohere_key, groq_key)

        # ---- Parse resumes (PDF / DOCX / TXT) ----
        docs = []
        processed_files = 0
        skipped = []

        for file_data in resume_files_data:
            filename = file_data["name"]
            file_content = file_data["content"]
            try:
                is_valid, msg = self.validate_file(file_content, filename)
                if not is_valid:
                    print(f"Skipping {filename}: {msg}")
                    skipped.append({"file": filename, "reason": msg})
                    continue

                parsed_docs = self.parse_resume(file_content, filename)
                if parsed_docs:
                    docs.extend(parsed_docs)
                    processed_files += 1
                    print(f"Processed: {filename}")
                else:
                    skipped.append({"file": filename, "reason": "No content extracted"})
            except Exception as e:
                print(f"Failed to process {filename}: {str(e)}")
                skipped.append({"file": filename, "reason": str(e)})

        if processed_files == 0:
            raise Exception("No valid resume files were processed!")

        # ---- Chunk ----
        splitter = RecursiveCharacterTextSplitter(chunk_size=1000, chunk_overlap=100)
        chunks = splitter.split_documents(docs)

        # ---- Embed + store ----
        self.create_vector_store(
            docs=chunks,
            collection=COLLECTION_NAME,
            qdrant_host=QDRANT_HOST,
            qdrant_api_key=QDRANT_API_KEY,
            dimension=DIMENSION,
            embeddings=embeddings
        )

        # ---- Retrieve top matching chunks (this score = NLP similarity) ----
        matches = self.query_resumes(
            query=job_desc,
            collection=COLLECTION_NAME,
            qdrant_host=QDRANT_HOST,
            qdrant_api_key=QDRANT_API_KEY,
            embeddings=embeddings,
            k=min(k * 5, len(chunks))
        )

        # ---- Group retrieved chunks by resume ----
        resume_matches = {}
        for match in matches:
            source = match["source"]
            resume_matches.setdefault(source, []).append(match)

        # ---- Score each resume ----
        results = []
        for resume_name, match_data in resume_matches.items():
            top_chunks = sorted(match_data, key=lambda x: x["score"], reverse=True)[:5]
            combined_text = " ".join(c["content"] for c in top_chunks)

            if len(combined_text) < 100:
                continue

            # NLP similarity score: average cosine similarity of top chunks, scaled to 0-100
            avg_similarity = sum(c["score"] for c in top_chunks) / len(top_chunks)
            similarity_score = round(max(0.0, min(1.0, avg_similarity)) * 100, 1)

            print(f"Analyzing: {resume_name}")
            analysis = self.analyze_resume_match(llm, job_desc, combined_text, resume_name)
            llm_score = max(0, min(100, analysis.get("llm_score", 50)))

            final_score = round(
                self.similarity_weight * similarity_score + self.llm_weight * llm_score, 1
            )

            skills = analysis.get("skills", [])
            if isinstance(skills, list):
                skills_str = ", ".join(skills) if skills else "Not specified"
            else:
                skills_str = str(skills)

            results.append({
                "Resume": resume_name,
                "Score": final_score,
                "NLP_Similarity_Score": similarity_score,
                "LLM_Assessment_Score": llm_score,
                "Skills": skills_str,
                "Experience": analysis.get("experience", "Not specified"),
                "Education": analysis.get("education", "Not specified"),
                "Strengths": analysis.get("strengths", "Not specified"),
                "Gaps": analysis.get("gaps", "Not specified"),
                "Overall_Assessment": analysis.get("overall_assessment", "Not specified"),
            })

        if not results:
            raise Exception("No results generated from resume analysis!")

        df = pd.DataFrame(results).sort_values("Score", ascending=False).reset_index(drop=True)
        df.insert(0, "Rank", range(1, len(df) + 1))

        if skipped:
            print(f"Skipped {len(skipped)} file(s): {skipped}")

        return df.head(k)

    # ------------------------------------------------------------------
    # Output helpers
    # ------------------------------------------------------------------
    def to_json(self, df: pd.DataFrame, job_desc: str = "") -> str:
        """Serialize ranked results to a JSON string with metadata."""
        payload = {
            "job_description": job_desc,
            "scoring_method": {
                "nlp_similarity_weight": self.similarity_weight,
                "llm_assessment_weight": self.llm_weight,
                "description": (
                    "Score = similarity_weight * NLP cosine-similarity score "
                    "+ llm_weight * LLM qualitative assessment score. "
                    "See SCORING_METHOD.md for full details."
                )
            },
            "total_candidates": len(df),
            "results": json.loads(df.to_json(orient="records"))
        }
        return json.dumps(payload, indent=2)
